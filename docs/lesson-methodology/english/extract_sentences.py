#!/usr/bin/env python
"""
Extract all sentences from Excel homework files into source_sentences.json.

Source: D:/GrammarMate/Грамматика до автоматизма/
Output: d:/Development/BaseGrammy/docs/lesson-methodology/english/source_sentences.json
"""

import json
import os
import re
import sys

import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = "D:/GrammarMate/Грамматика до автоматизма"
OUTPUT_PATH = "d:/Development/BaseGrammy/docs/lesson-methodology/english/source_sentences.json"

# Regex: rows that are just video markers like "2 видео", "3 видео"
VIDEO_MARKER_RE = re.compile(r'^\d+\s+видео$', re.IGNORECASE)

# Regex: section markers for Past Continuous (all-caps with colon)
SECTION_MARKER_RE = re.compile(r'^[А-ЯЁA-Z\s]+:.*$')


def convert_brackets(text):
    """Convert [hint] to (hint), preserving existing (context) markers."""
    if not text:
        return text
    # Replace all [ ... ] with ( ... )
    text = re.sub(r'\[([^\]]*)\]', r'(\1)', text)
    # Handle unclosed brackets like [Christmas (missing closing ])
    text = re.sub(r'\[([^\]\[]+)$', r'(\1)', text)
    return text


def is_skip_row(value):
    """Check if a row should be skipped."""
    if value is None:
        return True
    s = str(value).strip()
    if s == '':
        return True
    # Skip video markers like "2 видео"
    if VIDEO_MARKER_RE.match(s):
        return True
    return False


def detect_section_marker(value):
    """Check if a row is a section marker (all-caps text with colon)."""
    if value is None:
        return None
    s = str(value).strip()
    if SECTION_MARKER_RE.match(s):
        return s
    return None


def process_single_sheet(ws, sheet_name, topic, section=None):
    """Process a single worksheet and return list of sentence dicts."""
    sentences = []
    current_section = section

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=False):
        cell_value = row[0].value

        # Check for section markers first
        marker = detect_section_marker(cell_value)
        if marker:
            current_section = marker
            continue

        if is_skip_row(cell_value):
            continue

        text = str(cell_value).strip()
        text = convert_brackets(text)

        sentences.append({
            "ru": text,
            "source_sheet": sheet_name,
            "source_row": row[0].row,
            "section": current_section
        })

    return sentences


def process_file(filepath, topic, source_dir_name):
    """Process an Excel file and return list of sentences."""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    all_sentences = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]

        # Determine section for multi-sheet files
        section = None
        if topic == "perfect":
            # Map sheet names to readable sections
            section_map = {
                "Длитель": "Perfect Continuous (длительные)",
                "Соверш": "Perfect Simple (совершенные)",
                "Сов Практика": "Perfect Simple — Практика",
                "Сов+Длит": "Perfect Continuous + Simple"
            }
            section = section_map.get(sheet_name, sheet_name)
        elif topic == "articles":
            section = sheet_name

        sentences = process_single_sheet(ws, sheet_name, topic, section)
        all_sentences.extend(sentences)

    wb.close()
    return all_sentences


def main():
    os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)

    # Define all source files
    sources = [
        {
            "lesson_id": "L01_SENTENCES",
            "title": "L01 - Sentence Building",
            "topic": "sentence_building",
            "source_file": "1.Построение предложений",
            "filepath": os.path.join(BASE_DIR, "1.Построение предложений", "Текстовое ДЗ (построение предложений).xlsx"),
        },
        {
            "lesson_id": "L02_VERB_BE",
            "title": "L02 - Verb BE",
            "topic": "verb_be",
            "source_file": "2. Глагол BE",
            "filepath": os.path.join(BASE_DIR, "2. Глагол BE", "ДЗ-ТЕКСТ_ Глагол BE.xlsx"),
        },
        {
            "lesson_id": "L04_RELATIVE_CLAUSES",
            "title": "L04 - Relative Clauses (who/what/that)",
            "topic": "relative_clauses",
            "source_file": "4. Союзы WHO-WHAT-THAT",
            "filepath": os.path.join(BASE_DIR, "4. Союзы WHO-WHAT-THAT", "Союзы who-what-that.xlsx"),
        },
        {
            "lesson_id": "L05_VING",
            "title": "L05 - V-ing (Gerund)",
            "topic": "ving",
            "source_file": "5. Ving",
            "filepath": os.path.join(BASE_DIR, "5. Ving", "ДЗ-ТЕКСТ_ Ving.xlsx"),
        },
        {
            "lesson_id": "L06_ARTICLES",
            "title": "L06 - Articles",
            "topic": "articles",
            "source_file": "6. АРТИКЛИ",
            "filepath": os.path.join(BASE_DIR, "6. АРТИКЛИ", "6. Дз Текст АРТИКЛИ.xlsx"),
        },
        {
            "lesson_id": "L07_PRESENT_CONTINUOUS",
            "title": "L07 - Present Continuous",
            "topic": "present_continuous",
            "source_file": "7.Present Continuous",
            "filepath": os.path.join(BASE_DIR, "7.Present Continuous", "ДЗ-ТЕКСТ_ Present Continuous.xlsx"),
        },
        {
            "lesson_id": "L08_PAST_CONTINUOUS",
            "title": "L08 - Past Continuous",
            "topic": "past_continuous",
            "source_file": "8. Past Continuous",
            "filepath": os.path.join(BASE_DIR, "8. Past Continuous", "ДЗ-ТЕКСТ_ Past Continuous.xlsx"),
        },
        {
            "lesson_id": "L09_PERFECT",
            "title": "L09 - Perfect Tenses",
            "topic": "perfect",
            "source_file": "9. PERFECT",
            "filepath": os.path.join(BASE_DIR, "9. PERFECT", "ДЗ текст Perfect ВСЕ.xlsx"),
        },
        {
            "lesson_id": "L10_PASSIVE",
            "title": "L10 - Passive Voice",
            "topic": "passive",
            "source_file": "10.ПАССИВ",
            "filepath": os.path.join(BASE_DIR, "10.ПАССИВ", "ДЗ-ТЕКСТ_ PASSIVE.xlsx"),
        },
        {
            "lesson_id": "L11_CONDITIONALS",
            "title": "L11 - Conditionals (Если бы)",
            "topic": "conditionals",
            "source_file": "11. ЕСЛИ БЫ",
            "filepath": os.path.join(BASE_DIR, "11. ЕСЛИ БЫ", "ДЗ - Если бы.xlsx"),
        },
    ]

    lessons = []
    grand_total = 0

    for src in sources:
        print(f"\nProcessing: {src['source_file']} ...")
        if not os.path.exists(src['filepath']):
            print(f"  WARNING: File not found: {src['filepath']}")
            continue

        sentences = process_file(src['filepath'], src['topic'], src['source_file'])

        lesson = {
            "lesson_id": src['lesson_id'],
            "title": src['title'],
            "topic": src['topic'],
            "source_file": src['source_file'],
            "sentences": sentences
        }
        lessons.append(lesson)

        count = len(sentences)
        grand_total += count

        # Show section breakdown if applicable
        if src['topic'] in ('past_continuous', 'articles', 'perfect'):
            sections = {}
            for s in sentences:
                sec = s['section'] or '(no section)'
                sections[sec] = sections.get(sec, 0) + 1
            for sec, cnt in sections.items():
                print(f"  Section '{sec}': {cnt} sentences")
        else:
            print(f"  {count} sentences")

        # Print first 3 sentences as sample
        for s in sentences[:3]:
            sec_label = f" [{s['section']}]" if s['section'] else ""
            print(f"    Sample: {s['ru'][:80]}{sec_label}")

    result = {"lessons": lessons}

    with open(OUTPUT_PATH, 'w', encoding='utf-8') as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    print(f"\n{'='*60}")
    print(f"STATISTICS")
    print(f"{'='*60}")
    for lesson in lessons:
        print(f"  {lesson['lesson_id']:30s} ({lesson['topic']:20s}): {len(lesson['sentences']):4d} sentences")
    print(f"{'='*60}")
    print(f"  {'GRAND TOTAL':30s} {'':20s}: {grand_total:4d} sentences")
    print(f"{'='*60}")
    print(f"\nSaved to: {OUTPUT_PATH}")


if __name__ == '__main__':
    main()
